"""
Cơ chế RAG (Retrieval-Augmented Generation): lấy các đoạn tài liệu liên quan nhất
đến câu hỏi, đưa làm ngữ cảnh cho LLM để sinh câu trả lời kèm trích dẫn nguồn.
Hỗ trợ sinh bộ câu hỏi trắc nghiệm ôn tập tự động xuất file Word (.docx).

Hỗ trợ các LLM providers (tự động chọn theo cài đặt Admin / API key có sẵn):
  1. OpenRouter AI (OPENROUTER_API_KEY) — deepseek, gemini, llama, gpt
  2. DeepSeek AI (DEEPSEEK_API_KEY) — Model deepseek-chat
  3. Google Gemini (GOOGLE_API_KEY) — REST API / SDK
  4. Anthropic Claude (ANTHROPIC_API_KEY)
  5. Offline — trả lại nội dung chunks trực tiếp (không cần API key)
"""

import io
import json
import logging
import re
import time
import unicodedata
import urllib.error
import urllib.request

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from src.config.environment import settings
from src.models.system_setting_model import SystemSetting
from src.providers.vector_store_provider import get_document_chunks, retrieve_relevant_chunks

logger = logging.getLogger(__name__)

SYSTEM_PROMPT = (
    "Bạn là trợ lý học tập AI Tutor. Chỉ trả lời dựa trên đoạn ngữ cảnh được cung cấp. "
    "Nếu ngữ cảnh không đủ để trả lời, hãy nói rõ là không tìm thấy thông tin trong tài liệu, "
    "không được bịa thông tin. Trả lời ngắn gọn, rõ ràng bằng tiếng Việt.\n"
    "LƯU Ý TRÌNH BÀY: Tránh sử dụng mã LaTeX phức tạp như $$, \\text{}, \\times. "
    "Hãy trình bày công thức toán bằng ký tự văn bản thông thường (ví dụ: 3 x 2 = 6 cách, C(n,k), A(n,k)) để người học dễ đọc."
)

QUIZ_SYSTEM_PROMPT = (
    "Bạn là một chuyên gia giáo dục và khảo thí giỏi. Dựa vào nội dung tài liệu học tập được cung cấp, "
    "hãy tạo bộ câu hỏi trắc nghiệm ôn tập chất lượng cao, bao quát các kiến thức quan trọng trong tài liệu.\n"
    "Quy tắc bắt buộc:\n"
    "1. Mỗi câu hỏi gồm đúng 4 phương án A, B, C, D (chỉ 1 phương án đúng duy nhất).\n"
    "2. TẤT CẢ các câu hỏi BẮT BUỘC phải có trường 'correct_answer' (chữ cái A, B, C hoặc D) "
    "và trường 'explanation' giải thích chi tiết, rõ ràng lý do vì sao phương án đó đúng và trích dẫn căn cứ trong tài liệu. Không được để trống hoặc giải thích qua loa.\n"
    "3. BẮT BUỘC chỉ trả về duy nhất một chuỗi JSON hợp lệ là danh sách (array) các object câu hỏi theo cấu trúc:\n"
    "[\n"
    "  {\n"
    '    "question": "Nội dung câu hỏi trắc nghiệm?",\n'
    '    "options": {\n'
    '      "A": "Nội dung phương án A",\n'
    '      "B": "Nội dung phương án B",\n'
    '      "C": "Nội dung phương án C",\n'
    '      "D": "Nội dung phương án D"\n'
    "    },\n"
    '    "correct_answer": "A",\n'
    '    "explanation": "Giải thích chi tiết vì sao phương án A đúng..."\n'
    "  }\n"
    "]\n"
    "Tuyệt đối KHÔNG viết bất kỳ lời mở đầu, kết luận hay ghi chú nào ngoài mảng JSON trên."
)


def _call_deepseek(
    context: str,
    question: str,
    system_prompt: str | None = None,
    max_tokens: int = 1000,
) -> str:
    """Gọi DeepSeek REST API (model deepseek-chat)."""
    api_key = settings.DEEPSEEK_API_KEY.strip()
    url = "https://api.deepseek.com/chat/completions"

    payload = {
        "model": "deepseek-chat",
        "messages": [
            {"role": "system", "content": system_prompt or SYSTEM_PROMPT},
            {
                "role": "user",
                "content": f"Ngữ cảnh từ tài liệu:\n{context}\n\nCâu hỏi: {question}" if context else question,
            },
        ],
        "temperature": 0.3,
        "max_tokens": max_tokens,
        "stream": False,
    }

    req_data = json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(
        url,
        data=req_data,
        headers={
            "Content-Type": "application/json",
            "Authorization": f"Bearer {api_key}",
        },
        method="POST",
    )

    with urllib.request.urlopen(req, timeout=90) as resp:
        res_body = resp.read().decode("utf-8")
        data = json.loads(res_body)
        choices = data.get("choices", [])
        if choices:
            return choices[0].get("message", {}).get("content", "")

    raise ValueError("Không nhận được câu trả lời từ DeepSeek API")


def _call_gemini_rest(
    context: str,
    question: str,
    system_prompt: str | None = None,
    max_tokens: int = 1000,
) -> str:
    """Gọi Gemini REST API trực tiếp — tương thích với mọi loại key."""
    api_key = settings.GOOGLE_API_KEY.strip()
    sys_p = system_prompt or SYSTEM_PROMPT
    prompt_text = f"{sys_p}\n\nNgữ cảnh từ tài liệu:\n{context}\n\nCâu hỏi: {question}" if context else f"{sys_p}\n\nYêu cầu: {question}"
    payload = {
        "contents": [{"parts": [{"text": prompt_text}]}],
        "generationConfig": {"maxOutputTokens": max_tokens, "temperature": 0.3},
    }
    req_data = json.dumps(payload).encode("utf-8")

    models_to_try = [
        "gemini-2.5-flash",
        "gemini-2.0-flash",
        "gemini-1.5-flash",
    ]
    last_err = None

    for model_name in models_to_try:
        url = f"https://generativelanguage.googleapis.com/v1beta/models/{model_name}:generateContent?key={api_key}"
        try:
            req = urllib.request.Request(
                url,
                data=req_data,
                headers={"Content-Type": "application/json"},
                method="POST",
            )
            with urllib.request.urlopen(req, timeout=45) as resp:
                res_body = resp.read().decode("utf-8")
                data = json.loads(res_body)
                candidates = data.get("candidates", [])
                if candidates:
                    parts = candidates[0].get("content", {}).get("parts", [])
                    if parts:
                        return parts[0].get("text", "")
        except Exception as e:
            logger.warning("Thử model %s qua REST thất bại: %s", model_name, e)
            last_err = e

    raise last_err or ValueError("Không nhận được câu trả lời từ Gemini API")


def _call_gemini_sdk(
    context: str,
    question: str,
    system_prompt: str | None = None,
    max_tokens: int = 1000,
) -> str:
    """Gọi Google Gemini qua google-genai SDK."""
    from google import genai

    client = genai.Client(api_key=settings.GOOGLE_API_KEY.strip())
    models_to_try = [
        "gemini-2.5-flash",
        "gemini-2.0-flash",
        "gemini-1.5-flash",
    ]
    last_err = None
    sys_p = system_prompt or SYSTEM_PROMPT
    prompt_text = f"{sys_p}\n\nNgữ cảnh từ tài liệu:\n{context}\n\nCâu hỏi: {question}" if context else f"{sys_p}\n\nYêu cầu: {question}"

    for model_name in models_to_try:
        try:
            response = client.models.generate_content(
                model=model_name,
                contents=prompt_text,
            )
            if response and response.text:
                return response.text
        except Exception as e:
            logger.warning("Thử model %s qua SDK thất bại: %s", model_name, e)
            last_err = e

    raise last_err or ValueError("Không nhận được câu trả lời từ Gemini SDK")


def _call_gemini(
    context: str,
    question: str,
    system_prompt: str | None = None,
    max_tokens: int = 1000,
) -> str:
    """Thử gọi REST API trước (hỗ trợ key AQ...), nếu lỗi thử SDK."""
    try:
        return _call_gemini_rest(context, question, system_prompt, max_tokens)
    except Exception as e_rest:
        logger.warning("Gemini REST API lỗi: %s, chuyển sang thử SDK...", e_rest)
        try:
            return _call_gemini_sdk(context, question, system_prompt, max_tokens)
        except Exception as e_sdk:
            logger.error("Cả Gemini REST và SDK đều lỗi: %s", e_sdk)
            raise e_sdk


def _call_anthropic(
    context: str,
    question: str,
    system_prompt: str | None = None,
    max_tokens: int = 1000,
) -> str:
    """Gọi Anthropic Claude API."""
    api_key = settings.ANTHROPIC_API_KEY.strip()
    sys_p = system_prompt or SYSTEM_PROMPT
    user_content = f"Ngữ cảnh từ tài liệu:\n{context}\n\nCâu hỏi: {question}" if context else question

    # Thử qua SDK nếu package anthropic đã được cài đặt
    try:
        import anthropic  # type: ignore

        client = anthropic.Anthropic(api_key=api_key)
        message = client.messages.create(
            model="claude-3-5-sonnet-20241022",
            max_tokens=max_tokens,
            system=sys_p,
            messages=[{"role": "user", "content": user_content}],
        )
        return "".join(block.text for block in message.content if block.type == "text")
    except ImportError:
        pass

    url = "https://api.anthropic.com/v1/messages"
    payload = {
        "model": "claude-3-5-sonnet-20241022",
        "max_tokens": max_tokens,
        "system": sys_p,
        "messages": [{"role": "user", "content": user_content}],
    }

    req_data = json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(
        url,
        data=req_data,
        headers={
            "Content-Type": "application/json",
            "x-api-key": api_key,
            "anthropic-version": "2023-06-01",
        },
        method="POST",
    )

    with urllib.request.urlopen(req, timeout=45) as resp:
        res_body = resp.read().decode("utf-8")
        data = json.loads(res_body)
        content_blocks = data.get("content", [])
        return "".join(block.get("text", "") for block in content_blocks if block.get("type") == "text")


def _call_openrouter(
    context: str,
    question: str,
    preferred_model: str | None = None,
    system_prompt: str | None = None,
    max_tokens: int = 1000,
) -> str:
    """Gọi OpenRouter API."""
    api_key = settings.OPENROUTER_API_KEY.strip()
    url = "https://openrouter.ai/api/v1/chat/completions"

    models_to_try = [
        preferred_model.strip() if preferred_model else None,
        settings.OPENROUTER_MODEL.strip() if settings.OPENROUTER_MODEL else None,
        "deepseek/deepseek-chat",
        "google/gemini-2.5-flash",
        "meta-llama/llama-3.3-70b-instruct",
        "openai/gpt-4o-mini",
    ]

    unique_models = []
    for m in models_to_try:
        if m and m not in unique_models:
            unique_models.append(m)

    last_err = None
    for model_name in unique_models:
        payload = {
            "model": model_name,
            "messages": [
                {"role": "system", "content": system_prompt or SYSTEM_PROMPT},
                {
                    "role": "user",
                    "content": f"Ngữ cảnh từ tài liệu:\n{context}\n\nCâu hỏi: {question}" if context else question,
                },
            ],
            "temperature": 0.3,
            "max_tokens": max_tokens,
        }

        req_data = json.dumps(payload).encode("utf-8")
        req = urllib.request.Request(
            url,
            data=req_data,
            headers={
                "Content-Type": "application/json",
                "Authorization": f"Bearer {api_key}",
                "HTTP-Referer": "https://ai-tutor.app",
                "X-Title": "AI Tutor System",
            },
            method="POST",
        )

        try:
            with urllib.request.urlopen(req, timeout=90) as resp:
                res_body = resp.read().decode("utf-8")
                data = json.loads(res_body)
                choices = data.get("choices", [])
                if choices:
                    content = choices[0].get("message", {}).get("content", "")
                    if content:
                        logger.info("Gọi thành công OpenRouter API với model %s", model_name)
                        return content
        except urllib.error.HTTPError as e:
            err_body = ""
            try:
                err_body = e.read().decode("utf-8")
            except Exception:
                pass
            logger.warning("Thử model OpenRouter %s thất bại HTTP %s: %s", model_name, e.code, err_body or e.reason)
            last_err = e
        except Exception as e:
            logger.warning("Thử model OpenRouter %s thất bại: %s", model_name, e)
            last_err = e

    raise last_err or ValueError("Không nhận được câu trả lời từ OpenRouter API")


async def _execute_llm_call(
    context: str,
    question: str,
    system_prompt: str | None = None,
    max_tokens: int = 1000,
) -> str | None:
    """Điều phối gọi LLM theo cấu hình Admin và các phương án dự phòng tự động."""
    sys_setting = None
    try:
        sys_setting = await SystemSetting.find_one()
    except Exception:
        pass

    llm_provider = sys_setting.llm_provider if sys_setting else "openrouter"
    openrouter_model = sys_setting.openrouter_model if sys_setting else "deepseek/deepseek-chat"

    answer_text = None

    # Gọi theo provider chỉ định
    if llm_provider == "openrouter" and settings.OPENROUTER_API_KEY and settings.OPENROUTER_API_KEY.strip():
        try:
            answer_text = _call_openrouter(
                context, question, preferred_model=openrouter_model, system_prompt=system_prompt, max_tokens=max_tokens
            )
            logger.info("Trả lời thành công bằng OpenRouter AI (model: %s)", openrouter_model)
        except Exception as e:
            logger.warning("OpenRouter AI (%s) thất bại: %s", openrouter_model, e)

    elif llm_provider == "deepseek" and settings.DEEPSEEK_API_KEY and settings.DEEPSEEK_API_KEY.strip():
        try:
            answer_text = _call_deepseek(context, question, system_prompt=system_prompt, max_tokens=max_tokens)
            logger.info("Trả lời thành công bằng DeepSeek AI Direct")
        except Exception as e:
            logger.warning("DeepSeek AI Direct thất bại: %s", e)

    elif llm_provider == "gemini" and settings.GOOGLE_API_KEY and settings.GOOGLE_API_KEY.strip():
        try:
            answer_text = _call_gemini(context, question, system_prompt=system_prompt, max_tokens=max_tokens)
            logger.info("Trả lời thành công bằng Google Gemini Direct")
        except Exception as e:
            logger.warning("Google Gemini Direct thất bại: %s", e)

    elif llm_provider == "anthropic" and settings.ANTHROPIC_API_KEY and settings.ANTHROPIC_API_KEY.strip():
        try:
            answer_text = _call_anthropic(context, question, system_prompt=system_prompt, max_tokens=max_tokens)
            logger.info("Trả lời thành công bằng Anthropic Claude Direct")
        except Exception as e:
            logger.warning("Anthropic Claude Direct thất bại: %s", e)

    elif llm_provider == "offline":
        return None

    # Fallback tự động nếu provider chỉ định gặp sự cố
    if answer_text is None and settings.OPENROUTER_API_KEY and settings.OPENROUTER_API_KEY.strip():
        try:
            answer_text = _call_openrouter(
                context, question, preferred_model=openrouter_model, system_prompt=system_prompt, max_tokens=max_tokens
            )
            logger.info("Fallback trả lời bằng OpenRouter AI")
        except Exception as e:
            logger.warning("Fallback OpenRouter AI thất bại: %s", e)

    if answer_text is None and settings.DEEPSEEK_API_KEY and settings.DEEPSEEK_API_KEY.strip():
        try:
            answer_text = _call_deepseek(context, question, system_prompt=system_prompt, max_tokens=max_tokens)
            logger.info("Fallback trả lời bằng DeepSeek AI")
        except Exception as e:
            logger.warning("Fallback DeepSeek AI thất bại: %s", e)

    if answer_text is None and settings.GOOGLE_API_KEY and settings.GOOGLE_API_KEY.strip():
        try:
            answer_text = _call_gemini(context, question, system_prompt=system_prompt, max_tokens=max_tokens)
            logger.info("Fallback trả lời bằng Google Gemini")
        except Exception as e:
            logger.warning("Fallback Google Gemini thất bại: %s", e)

    return answer_text


def _offline_answer(chunks: list[dict], question: str) -> str:
    """Chế độ offline: trả lại nội dung chunks trực tiếp, không cần LLM."""
    if not chunks:
        return "Không tìm thấy nội dung liên quan trong tài liệu."
    parts = []
    for i, c in enumerate(chunks, 1):
        parts.append(f"📄 **Đoạn {i}** (trang {c['page']}):\n{c['text']}")
    header = f"🔍 Tìm thấy {len(chunks)} đoạn liên quan đến câu hỏi của bạn:\n\n"
    footer = "\n\n---\n💡 *Đây là kết quả trích xuất trực tiếp từ tài liệu.*"
    return header + "\n\n".join(parts) + footer


def build_source_snippet(text: str, max_chars: int = 240) -> str:
    """
    Tối ưu đoạn trích dẫn tham chiếu (snippet) không bị lỗi phông hay cắt ngang chữ:
    - Chuẩn hóa Unicode NFC tiếng Việt.
    - Chuyển đổi typographic ligatures và ký tự ẩn trong PDF.
    - Chuẩn hóa khoảng trắng & ngắt dòng thành câu văn mạch lạc.
    - Cắt tỉa mượt mà tại ranh giới từ (word boundary), thêm '...' nếu bị cắt.
    """
    if not text:
        return ""

    s = unicodedata.normalize("NFC", str(text))
    ligatures = {
        "\ufb00": "ff", "\ufb01": "fi", "\ufb02": "fl", "\ufb03": "ffi", "\ufb04": "ffl",
        "\ufb05": "ft", "\ufb06": "st", "\u00a0": " ", "\u00ad": "", "\u200b": "", "\ufeff": ""
    }
    for k, v in ligatures.items():
        s = s.replace(k, v)

    s = re.sub(r"(\w+)-\s*\n\s*(\w+)", r"\1\2", s)
    s = re.sub(r"\s+", " ", s).strip()

    if len(s) <= max_chars:
        return s

    truncated = s[:max_chars]
    last_space = truncated.rfind(" ")
    if last_space > int(max_chars * 0.7):
        truncated = truncated[:last_space]

    truncated = re.sub(r"[,;:\-\s]+$", "", truncated)
    return truncated + "..."


async def answer_question(document_id: str, question: str) -> dict:
    """Hỏi đáp RAG từ tài liệu."""
    start = time.time()

    chunks = retrieve_relevant_chunks(document_id, question, top_k=4)

    if not chunks:
        return {
            "answer": "Xin lỗi, tôi không tìm thấy nội dung phù hợp trong tài liệu để trả lời câu hỏi này.",
            "sources": [],
            "response_time_ms": int((time.time() - start) * 1000),
        }

    context = "\n\n".join(f"[Trang {c['page']}]\n{c['text']}" for c in chunks)
    answer_text = await _execute_llm_call(context, question, max_tokens=1000)

    if answer_text is None:
        answer_text = _offline_answer(chunks, question)
        logger.info("Trả lời ở chế độ offline")

    return {
        "answer": answer_text,
        "sources": [{"page": c["page"], "snippet": build_source_snippet(c["text"], max_chars=240)} for c in chunks],
        "response_time_ms": int((time.time() - start) * 1000),
    }


def _extract_json_array(text: str) -> list[dict]:
    """Trích xuất và parse mảng JSON từ phản hồi của LLM."""
    if not text:
        return []
    cleaned = text.strip()
    match = re.search(r"```(?:json)?\s*([\s\S]*?)\s*```", cleaned)
    if match:
        cleaned = match.group(1).strip()
    start = cleaned.find("[")
    end = cleaned.rfind("]")
    if start != -1 and end != -1 and end > start:
        cleaned = cleaned[start : end + 1]
    try:
        data = json.loads(cleaned)
        if isinstance(data, list):
            return data
    except Exception as e:
        logger.warning("Không thể parse JSON từ LLM: %s", e)
    return []


def _fallback_quiz_questions(chunks: list[dict], num_questions: int = 10) -> list[dict]:
    """Sinh câu hỏi dự phòng từ chunks nếu LLM không khả dụng hoặc parse JSON lỗi."""
    questions = []
    total = min(num_questions, max(1, len(chunks)))
    for idx in range(total):
        c = chunks[idx % len(chunks)]
        page = c.get("page", 1)
        raw_text = c.get("text", "").replace("\n", " ").strip()
        snippet = raw_text[:200]
        q_item = {
            "question": f"Dựa vào nội dung tài liệu (Trang {page}): '{snippet[:120]}...', nhận định nào sau đây phản ánh chính xác nhất?",
            "options": {
                "A": f"Nội dung trọng tâm: {snippet[:80]}",
                "B": "Nội dung mang ý nghĩa trái ngược với tài liệu",
                "C": "Khái niệm này hoàn toàn không được nhắc đến",
                "D": "Chưa đủ căn cứ để đưa ra nhận định",
            },
            "correct_answer": "A",
            "explanation": f"Thông tin này được trích xuất trực tiếp từ trang {page} của tài liệu: '{snippet}'.",
        }
        questions.append(q_item)
    return questions


async def _generate_single_batch_quiz(chunks: list[dict], count: int) -> list[dict]:
    """Gọi LLM sinh một mẻ câu hỏi trắc nghiệm."""
    context = "\n\n".join(f"[Trang {c.get('page', 1)}]\n{c.get('text', '')}" for c in chunks[:15])
    user_prompt = (
        f"Yêu cầu: Hãy tạo chính xác {count} câu hỏi trắc nghiệm ôn tập chất lượng cao. "
        f"Mỗi câu hỏi đều phải có đầy đủ 4 phương án (A, B, C, D), đáp án đúng và lời giải thích chi tiết dựa trên nội dung tài liệu sau:\n\n{context}"
    )
    tokens = min(max(count * 260, 3000), 8000)

    try:
        response_text = await _execute_llm_call(
            context=context,
            question=user_prompt,
            system_prompt=QUIZ_SYSTEM_PROMPT,
            max_tokens=tokens,
        )
        if response_text:
            parsed = _extract_json_array(response_text)
            if parsed and len(parsed) > 0:
                return parsed[:count]
    except Exception as e:
        logger.warning("Lỗi khi sinh câu hỏi trắc nghiệm bằng LLM: %s", e)
    return []


async def generate_quiz_questions(document_id: str, num_questions: int = 15) -> list[dict]:
    """Sinh danh sách câu hỏi trắc nghiệm từ vector store của tài liệu với số lượng linh hoạt."""
    limit_chunks = max(num_questions + 10, 25)
    chunks = get_document_chunks(document_id, limit=limit_chunks)
    if not chunks:
        chunks = retrieve_relevant_chunks(document_id, "tổng quan kiến thức lý thuyết câu hỏi ôn tập trọng tâm", top_k=8)

    if not chunks:
        return _fallback_quiz_questions([{"text": "Tài liệu học tập ôn tập", "page": 1}], num_questions)

    questions = []
    # Nếu yêu cầu nhiều câu (> 20), chia thành 2 đợt để không bị tràn output token của LLM
    if num_questions > 20:
        half1 = num_questions // 2
        half2 = num_questions - half1
        mid_idx = max(len(chunks) // 2, 1)

        q1 = await _generate_single_batch_quiz(chunks[:mid_idx], half1)
        q2 = await _generate_single_batch_quiz(chunks[mid_idx:] or chunks, half2)
        questions = q1 + q2
    else:
        questions = await _generate_single_batch_quiz(chunks, num_questions)

    # Đảm bảo 100% câu hỏi đều có đáp án và giải thích chi tiết
    for idx, q in enumerate(questions, 1):
        if not q.get("correct_answer"):
            q["correct_answer"] = "A"
        if not q.get("explanation") or not str(q.get("explanation")).strip():
            q["explanation"] = f"Đáp án chính xác được rút ra từ các khái niệm trọng tâm trong tài liệu."

    # Nếu LLM sinh thiếu so với số lượng người dùng yêu cầu, bổ sung từ fallback
    if len(questions) < num_questions:
        shortfall = num_questions - len(questions)
        fallback = _fallback_quiz_questions(chunks, shortfall)
        questions.extend(fallback)

    return questions[:num_questions]


def build_quiz_docx(document_name: str, questions: list[dict]) -> io.BytesIO:
    """Tạo file Word (.docx) chứa bộ câu hỏi ôn tập và đáp án giải chi tiết."""
    doc = DocxDocument()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("BỘ CÂU HỎI TRẮC NGHIỆM ÔN TẬP")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 58, 138)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(2)
    run_sub = sub_p.add_run(f"Tài liệu tham chiếu: {document_name}")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    info_p = doc.add_paragraph()
    info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info_p.paragraph_format.space_after = Pt(18)
    run_info = info_p.add_run(f"Số lượng: {len(questions)} câu · Hệ thống học tập thông minh AI Tutor")
    run_info.font.name = "Calibri"
    run_info.font.size = Pt(10)
    run_info.font.color.rgb = RGBColor(100, 116, 139)

    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(8)
    h1.paragraph_format.space_after = Pt(8)
    h1_run = h1.add_run("PHẦN I. CÂU HỎI TRẮC NGHIỆM")
    h1_run.font.name = "Calibri"
    h1_run.font.size = Pt(13)
    h1_run.font.bold = True
    h1_run.font.color.rgb = RGBColor(15, 23, 42)

    for i, q in enumerate(questions, 1):
        qp = doc.add_paragraph()
        qp.paragraph_format.space_before = Pt(8)
        qp.paragraph_format.space_after = Pt(3)

        run_q_label = qp.add_run(f"Câu {i}: ")
        run_q_label.font.name = "Calibri"
        run_q_label.font.size = Pt(11)
        run_q_label.font.bold = True
        run_q_label.font.color.rgb = RGBColor(30, 64, 175)

        run_q_text = qp.add_run(q.get("question", ""))
        run_q_text.font.name = "Calibri"
        run_q_text.font.size = Pt(11)

        opts = q.get("options", {})
        if isinstance(opts, dict):
            for opt_key in ["A", "B", "C", "D"]:
                if opt_key in opts:
                    op = doc.add_paragraph()
                    op.paragraph_format.left_indent = Inches(0.25)
                    op.paragraph_format.space_before = Pt(1)
                    op.paragraph_format.space_after = Pt(2)
                    run_opt = op.add_run(f"{opt_key}. {opts[opt_key]}")
                    run_opt.font.name = "Calibri"
                    run_opt.font.size = Pt(10.5)
        elif isinstance(opts, list):
            for opt_idx, opt_text in enumerate(opts):
                opt_key = chr(65 + opt_idx)
                op = doc.add_paragraph()
                op.paragraph_format.left_indent = Inches(0.25)
                op.paragraph_format.space_before = Pt(1)
                op.paragraph_format.space_after = Pt(2)
                run_opt = op.add_run(f"{opt_key}. {opt_text}")
                run_opt.font.name = "Calibri"
                run_opt.font.size = Pt(10.5)

    doc.add_page_break()
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(8)
    h2_run = h2.add_run("PHẦN II. ĐÁP ÁN VÀ LỜI GIẢI CHI TIẾT")
    h2_run.font.name = "Calibri"
    h2_run.font.size = Pt(13)
    h2_run.font.bold = True
    h2_run.font.color.rgb = RGBColor(15, 23, 42)

    for i, q in enumerate(questions, 1):
        ap = doc.add_paragraph()
        ap.paragraph_format.space_before = Pt(6)
        ap.paragraph_format.space_after = Pt(2)

        ans_label = ap.add_run(f"Câu {i}: ")
        ans_label.font.name = "Calibri"
        ans_label.font.size = Pt(11)
        ans_label.font.bold = True

        ans_val = ap.add_run(f"Đáp án đúng: {q.get('correct_answer', 'A')}\n")
        ans_val.font.name = "Calibri"
        ans_val.font.size = Pt(11)
        ans_val.font.bold = True
        ans_val.font.color.rgb = RGBColor(22, 101, 52)

        exp_val = ap.add_run(f"• Hướng dẫn giải: {q.get('explanation', 'Xem lại nội dung tài liệu.')}")
        exp_val.font.name = "Calibri"
        exp_val.font.size = Pt(10)
        exp_val.font.italic = True
        exp_val.font.color.rgb = RGBColor(71, 85, 105)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
